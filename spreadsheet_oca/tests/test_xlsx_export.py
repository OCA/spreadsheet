# Copyright 2025 Ledo Enterprises LLC
# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl).
"""
Tests for headless XLSX export (spreadsheet_xlsx_export).
"""

import base64
import io

import openpyxl

from odoo.tests import TransactionCase
from odoo.tools import mute_logger

from ..models.spreadsheet_xlsx_export import (
    SpreadsheetXlsxExporter,
    _coerce_value,
    _format_measure_name,
)


class TestHelpers(TransactionCase):
    """Unit tests for pure helper functions."""

    def test_coerce_value_int(self):
        self.assertEqual(_coerce_value("42"), 42)

    def test_coerce_value_float(self):
        self.assertAlmostEqual(_coerce_value("3.14"), 3.14)

    def test_coerce_value_float_comma(self):
        # Commas stripped as thousands separators: "1,234.56" → 1234.56
        self.assertAlmostEqual(_coerce_value("1,234.56"), 1234.56)

    def test_coerce_value_string(self):
        self.assertEqual(_coerce_value("hello"), "hello")

    def test_coerce_value_bool_passthrough(self):
        self.assertTrue(_coerce_value(True))
        self.assertFalse(_coerce_value(False))

    def test_format_measure_name_with_colon(self):
        self.assertEqual(
            _format_measure_name("amount_total:sum"),
            "Amount Total (Sum)",
        )

    def test_format_measure_name_plain(self):
        self.assertEqual(_format_measure_name("amount_total"), "Amount Total")


class TestXlsxExport(TransactionCase):
    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.spreadsheet = cls.env["spreadsheet.spreadsheet"].create(
            {"name": "Test Export Sheet"}
        )

    def _load_wb(self, xlsx_bytes):
        """Open xlsx bytes as an openpyxl workbook."""
        return openpyxl.load_workbook(io.BytesIO(xlsx_bytes))

    # ── render() basics ───────────────────────────────────────────────────────

    def test_render_returns_bytes(self):
        xlsx = SpreadsheetXlsxExporter(self.env, self.spreadsheet).render()
        self.assertIsInstance(xlsx, bytes)
        self.assertGreater(len(xlsx), 0)

    def test_render_empty_spreadsheet_creates_fallback_sheet(self):
        self.spreadsheet.write({"spreadsheet_raw": {}})
        wb = self._load_wb(SpreadsheetXlsxExporter(self.env, self.spreadsheet).render())
        self.assertEqual(len(wb.sheetnames), 1)
        self.assertEqual(wb.sheetnames[0], "Empty")

    def test_render_static_sheet_values(self):
        raw = {
            "sheets": [
                {
                    "id": "s1",
                    "name": "Summary",
                    "cells": {
                        "A1": {"content": "Revenue"},
                        "B1": {"content": "12345"},
                        "A2": {"content": "Cost"},
                        "B2": {"content": "7000"},
                    },
                }
            ]
        }
        self.spreadsheet.write({"spreadsheet_raw": raw})
        wb = self._load_wb(SpreadsheetXlsxExporter(self.env, self.spreadsheet).render())

        self.assertIn("Summary", wb.sheetnames)
        ws = wb["Summary"]
        self.assertEqual(ws.cell(1, 1).value, "Revenue")
        self.assertEqual(ws.cell(1, 2).value, 12345)  # numeric coercion
        self.assertEqual(ws.cell(2, 1).value, "Cost")
        self.assertEqual(ws.cell(2, 2).value, 7000)

    def test_render_formula_cells_written_verbatim(self):
        raw = {
            "sheets": [
                {
                    "id": "s1",
                    "name": "Formulas",
                    "cells": {
                        "A1": {"content": "=SUM(B1:B10)"},
                    },
                }
            ]
        }
        self.spreadsheet.write({"spreadsheet_raw": raw})
        wb = self._load_wb(SpreadsheetXlsxExporter(self.env, self.spreadsheet).render())
        ws = wb["Formulas"]
        # Formula preserved as a string (server can't evaluate it)
        self.assertEqual(ws.cell(1, 1).value, "=SUM(B1:B10)")

    def test_render_multiple_static_sheets(self):
        raw = {
            "sheets": [
                {"id": "s1", "name": "Sheet1", "cells": {"A1": {"content": "A"}}},
                {"id": "s2", "name": "Sheet2", "cells": {"A1": {"content": "B"}}},
            ]
        }
        self.spreadsheet.write({"spreadsheet_raw": raw})
        wb = self._load_wb(SpreadsheetXlsxExporter(self.env, self.spreadsheet).render())
        self.assertEqual(wb.sheetnames, ["Sheet1", "Sheet2"])

    @mute_logger("odoo.addons.spreadsheet_oca.models.pivot_data")
    def test_render_with_unknown_pivot_model_graceful(self):
        """Unknown model in pivot definition should not crash; pivot is skipped."""
        raw = {
            "sheets": [{"id": "s1", "name": "Data", "cells": {}}],
            "pivots": {
                "1": {
                    "type": "ODOO",
                    "model": "nonexistent.model.xyz",
                    "domain": [],
                    "context": {},
                    "rows": [],
                    "columns": [],
                    "measures": [],
                    "name": "Bad Pivot",
                }
            },
        }
        self.spreadsheet.write({"spreadsheet_raw": raw})
        # Should not raise — unknown model is silently skipped
        wb = self._load_wb(SpreadsheetXlsxExporter(self.env, self.spreadsheet).render())
        # Only the static sheet, no pivot sheet for the invalid model
        self.assertNotIn("Bad Pivot", wb.sheetnames)
        self.assertIn("Data", wb.sheetnames)

    def test_render_pivot_non_odoo_type_skipped(self):
        """Non-ODOO pivot types are skipped (no extra worksheet)."""
        raw = {
            "sheets": [{"id": "s1", "name": "Data", "cells": {}}],
            "pivots": {
                "1": {"type": "CUSTOM", "name": "Custom Pivot"},
            },
        }
        self.spreadsheet.write({"spreadsheet_raw": raw})
        wb = self._load_wb(SpreadsheetXlsxExporter(self.env, self.spreadsheet).render())
        # Only the static sheet, no pivot sheet for CUSTOM type
        self.assertEqual(wb.sheetnames, ["Data"])

    def test_render_pivot_grand_total_only(self):
        """
        A pivot with no row/col dimensions produces a grand-total table.
        Use res.partner (always present) with count measure.
        """
        raw = {
            "sheets": [{"id": "s1", "name": "Data", "cells": {}}],
            "pivots": {
                "1": {
                    "type": "ODOO",
                    "model": "res.partner",
                    "domain": [["id", "=", 1]],  # narrow domain for speed
                    "context": {},
                    "rows": [],
                    "columns": [],
                    "measures": [],
                    "name": "Partner Count",
                }
            },
        }
        self.spreadsheet.write({"spreadsheet_raw": raw})
        wb = self._load_wb(SpreadsheetXlsxExporter(self.env, self.spreadsheet).render())
        self.assertIn("Partner Count", wb.sheetnames)

    def test_render_pivot_row_breakdown(self):
        """Row-only pivot creates a table with header + data rows."""
        raw = {
            "sheets": [{"id": "s1", "name": "Data", "cells": {}}],
            "pivots": {
                "1": {
                    "type": "ODOO",
                    "model": "res.partner",
                    "domain": [["id", "in", [1, 3]]],
                    "context": {},
                    "rows": [{"fieldName": "id", "order": "asc", "type": "integer"}],
                    "columns": [],
                    "measures": [],
                    "name": "By ID",
                }
            },
        }
        self.spreadsheet.write({"spreadsheet_raw": raw})
        # Should not raise; pivot sheet created
        wb = self._load_wb(SpreadsheetXlsxExporter(self.env, self.spreadsheet).render())
        self.assertIn("By ID", wb.sheetnames)

    def test_render_pivot_crosstab(self):
        """Cross-tab pivot (both row and column groupby) produces correct layout.

        Uses res.partner with row=is_company and col=type, exercising the
        _write_crosstab code path.
        """
        raw = {
            "sheets": [{"id": "s1", "name": "Data", "cells": {}}],
            "pivots": {
                "1": {
                    "type": "ODOO",
                    "model": "res.partner",
                    "domain": [],
                    "context": {},
                    "rows": [
                        {"fieldName": "is_company", "order": "asc", "type": "boolean"}
                    ],
                    "columns": [
                        {"fieldName": "type", "order": "asc", "type": "selection"}
                    ],
                    "measures": [],
                    "name": "Cross-Tab Test",
                }
            },
        }
        self.spreadsheet.write({"spreadsheet_raw": raw})
        wb = self._load_wb(SpreadsheetXlsxExporter(self.env, self.spreadsheet).render())
        self.assertIn("Cross-Tab Test", wb.sheetnames)
        ws = wb["Cross-Tab Test"]
        # Row 1 is the title row with the pivot name
        self.assertEqual(ws.cell(1, 1).value, "Cross-Tab Test")
        # Row 4+ should contain column headers — at minimum the row dimension
        # header ("is_company") should appear somewhere in the header area.
        header_values = [ws.cell(4, c).value for c in range(1, ws.max_column + 1)]
        self.assertIn("is_company", header_values)

    # ── action_export_xlsx ────────────────────────────────────────────────────

    def test_action_export_xlsx_returns_act_url(self):
        self.spreadsheet.write({"spreadsheet_raw": {}})
        action = self.spreadsheet.action_export_xlsx()
        self.assertEqual(action["type"], "ir.actions.act_url")
        self.assertIn("/web/content/", action["url"])
        self.assertIn("download=true", action["url"])

    def test_action_export_xlsx_creates_attachment(self):
        self.spreadsheet.write({"spreadsheet_raw": {}})
        before = self.env["ir.attachment"].search_count(
            [
                ("res_model", "=", "spreadsheet.spreadsheet"),
                ("res_id", "=", self.spreadsheet.id),
            ]
        )
        self.spreadsheet.action_export_xlsx()
        after = self.env["ir.attachment"].search_count(
            [
                ("res_model", "=", "spreadsheet.spreadsheet"),
                ("res_id", "=", self.spreadsheet.id),
            ]
        )
        self.assertEqual(after, before + 1)

    def test_action_export_xlsx_filename(self):
        self.spreadsheet.write({"spreadsheet_raw": {}, "name": "Sales KPI"})
        self.spreadsheet.action_export_xlsx()
        att = self.env["ir.attachment"].search(
            [
                ("res_model", "=", "spreadsheet.spreadsheet"),
                ("res_id", "=", self.spreadsheet.id),
            ],
            order="id desc",
            limit=1,
        )
        self.assertEqual(att.name, "Sales KPI.xlsx")

    def test_action_export_xlsx_attachment_is_valid_xlsx(self):
        self.spreadsheet.write({"spreadsheet_raw": {}})
        self.spreadsheet.action_export_xlsx()
        att = self.env["ir.attachment"].search(
            [
                ("res_model", "=", "spreadsheet.spreadsheet"),
                ("res_id", "=", self.spreadsheet.id),
            ],
            order="id desc",
            limit=1,
        )
        xlsx_bytes = base64.b64decode(att.datas)
        # Should open as a valid workbook
        wb = self._load_wb(xlsx_bytes)
        self.assertGreater(len(wb.sheetnames), 0)

    # ── get_xlsx_bytes ────────────────────────────────────────────────────────

    def test_get_xlsx_bytes_returns_base64(self):
        self.spreadsheet.write({"spreadsheet_raw": {}})
        b64 = self.env["spreadsheet.spreadsheet"].get_xlsx_bytes(self.spreadsheet.id)
        self.assertIsInstance(b64, str)
        # Must be valid base64
        decoded = base64.b64decode(b64)
        self._load_wb(decoded)  # valid xlsx


class TestXlsxPivotLabels(TransactionCase):
    """A many2one group must export as its name, not its database id."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.country_be = cls.env.ref("base.be")
        cls.country_us = cls.env.ref("base.us")
        cls.partners = cls.env["res.partner"].create(
            [
                {"name": "Xlsx Alpha", "country_id": cls.country_be.id},
                {"name": "Xlsx Beta", "country_id": cls.country_us.id},
            ]
        )
        cls.spreadsheet = cls.env["spreadsheet.spreadsheet"].create(
            {
                "name": "Pivot Export Sheet",
                "spreadsheet_raw": {
                    "sheets": [{"name": "Data", "cells": {}}],
                    "pivots": {
                        "1": {
                            "type": "ODOO",
                            "id": "1",
                            "name": "Partners by Country",
                            "model": "res.partner",
                            "domain": [("id", "in", cls.partners.ids)],
                            "context": {},
                            "measures": [{"id": "__count", "fieldName": "__count"}],
                            "rows": [{"fieldName": "country_id"}],
                            "columns": [],
                        }
                    },
                },
            }
        )

    def test_pivot_rows_export_country_names_not_ids(self):
        xlsx = SpreadsheetXlsxExporter(self.env, self.spreadsheet).render()
        wb = openpyxl.load_workbook(io.BytesIO(xlsx))

        written = set()
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                written.update(str(c) for c in row if c is not None)

        self.assertIn(self.country_be.display_name, written)
        self.assertIn(self.country_us.display_name, written)
        self.assertNotIn(str(self.country_be.id), written)
        self.assertNotIn(str(self.country_us.id), written)
